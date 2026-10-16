from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
sec.page_height = Inches(11.7)
sec.page_width = Inches(8.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.07
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(13)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('Level Movement study guide  |  ')
r.font.size = Pt(9)
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text, 2)
def qa(q,a):
    x=doc.add_paragraph(); x.add_run(q+' ').bold=True; x.add_run(a)
def page(title):
    doc.add_page_break(); doc.add_heading(title,1)
def source(text):
    x=doc.add_paragraph('Code to revisit: '+text); x.runs[0].italic=True; x.runs[0].font.size=Pt(9)

doc.add_heading('Level Movement programming principles',0)
doc.add_paragraph('Project explanation and intermediate interview study guide','Subtitle')
p('Learn to explain each design choice using four points: the problem, the implementation, the benefit, and the tradeoff. This guide connects those points to your vehicle management application and provides answers you can practise aloud.')
h('Your project in one minute')
p('“I built a vehicle management application with a React and TypeScript frontend and an ASP.NET Core backend. The backend separates HTTP handling, application workflows, domain concepts, and persistence. It uses MediatR commands and queries, validation behaviours, repositories, and a unit of work. EF Core maps value objects into SQL Server columns. Updates and soft deletes use optimistic concurrency so stale forms cannot silently overwrite newer changes. Booking validation checks whether a vehicle is eligible; it does not create a booking.”')
h('The main ideas you have used')
for t in [
'Separation of concerns and dependency inversion keep HTTP, business concepts, and database access in distinct places.',
'CQRS and a mediator give reads and writes explicit request types and shared processing steps.',
'Value objects centralise normalisation, equality, and rules for registration numbers, years, and rates.',
'Validation, unique database constraints, concurrency tokens, and soft deletion protect different aspects of data integrity.',
'React hooks separate interaction logic from page markup; query caching manages data received from the server.',
'Automated tests exercise domain rules, HTTP behaviour, persistence interactions, and frontend workflows.'
]: doc.add_paragraph(t,'List Bullet')
h('How to use this guide')
p('Read the architecture, request flow, domain, persistence, API, and React sections with the relevant code open. Review the DDD changes and tradeoffs, then use the intermediate concepts and practice questions to rehearse without looking at the answers.')
p('A principle is a general rule, such as separation of concerns. A pattern is a repeatable design, such as Repository. A tool is an implementation aid, such as EF Core or MediatR. Explain the problem solved before naming the tool.')
p('Project snapshot: 10 October 2026. File references are relative to the repository root. Examples describe the implementation in this snapshot; proposed improvements are identified separately.')

page('Architecture and SOLID')
h('Layered architecture and separation of concerns')
p('Domain contains the Vehicle aggregate root, value objects, and repository contracts. Application contains commands, queries, handlers, organisers, validators, DTOs, service contracts, and workflow services. Infrastructure implements EF persistence, mappings, concurrency stamping, and the unit of work. API contains controllers, exception handling, Swagger configuration, and dependency registration.')
p('Project references point inward: Application references Domain; Infrastructure references Application and Domain; API references Application and Infrastructure. Domain has no project or framework package references. The API acts as the composition root, where interfaces are connected to implementations.')
qa('Why do this?', 'A change to HTTP response formatting should not require changing registration-number rules. Separating reasons to change makes code easier to understand and test. The cost is more files and navigation for even a small feature.')
h('SOLID in this project')
qa('Single responsibility', 'VehicleConfig handles persistence mapping, RegistrationNumber handles its value rules, and VehiclesController handles HTTP entry points. Responsibility means a coherent reason to change, not necessarily one method per class.')
qa('Open closed', 'Pipeline behaviours provide a place to add shared processing without rewriting every handler. This supports extension, but does not mean existing code will never need modification.')
qa('Liskov substitution', 'An IVehicleRepository implementation should preserve its contract when substituted: missing lookups return null where declared, cancellation is respected, and writes retain concurrency semantics. Interfaces alone do not prove this principle is satisfied.')
qa('Interface segregation', 'IBookingEligibilityService and IUnitOfWork expose small, focused contracts. IVehicleRepository combines reads and writes; if clients need different subsets, it could be split. More interfaces are useful only when they clarify real boundaries.')
qa('Dependency inversion', 'Organisers depend on IVehiclesService and services depend on IVehicleRepository. The concrete database implementation is supplied externally. High-level workflows need not construct an EF repository themselves.')
h('Dependency injection and lifetimes')
p('ConfigurationExtension registers repositories, services, organisers, the interceptor, and the unit of work as scoped. AddDbContext also uses a scoped context by default. Within an HTTP request, these components can share the same context and tracked changes. Transient means a new instance for each resolution; singleton means one for the container lifetime. A singleton should not capture a scoped DbContext.')
source('backend/src/LevelMovement.Api/Extensions/ConfigurationExtension.cs; the four backend project files; Application/Interfaces and Domain/Interfaces under backend/src.')

page('CQRS mediator and request flow')
h('Commands change state and queries read state')
p('UpsertVehicleCommand and DeleteVehicleCommand implement ICommand<TResponse>. Read requests and ValidateBookingQuery implement IRequest<TResponse>. This is a lightweight CQRS design: command and query responsibilities are separated, but they use the same database and entity model. CQRS does not require event sourcing, separate databases, or a message broker.')
p('BookingsController uses HTTP POST for validation, but the application request is a query because it does not persist a change. HTTP verbs and application command/query classification describe different concerns.')
h('Trace a vehicle save')
for t in [
'The edit view submits values through useVehicleEditController and vehiclesApi.save to POST /api/vehicles.',
'ASP.NET binds JSON to UpsertVehicleCommand. VehiclesController calls ISender.Send.',
'RequestLoggingBehavior records the request type. ValidationBehavior runs the request validators. UnitOfWorkBehavior wraps command execution.',
'The handler calls UpsertVehicleOrganiser, which constructs value objects and adapts the existing upsert request to an explicit RegisterAsync or UpdateAsync application operation.',
'The Application VehiclesService calls Vehicle.Register for creation, or loads the existing vehicle and calls UpdateDetails. The repository stages the addition or version-checked update without deciding the business rules.',
'After the handler returns, UnitOfWorkBehavior saves. The interceptor assigns a new concurrency stamp and EF persists the change.',
'The controller calls ToResponse after the save, so the DTO includes the saved stamp, and returns 201 for creation or 200 for update.'
]: doc.add_paragraph(t,'List Number')
h('Why behaviours matter')
p('Validation and saving are cross-cutting concerns: many requests need the same processing. A behaviour wraps the next step in the pipeline. If validation throws, the handler is not called. If the handler throws, the save after next() is not reached. Queries do not implement ICommand, so they do not enter the command-only save behaviour.')
qa('Is MediatR a queue?', 'No. Here it dispatches requests within the running process. There is no durable delivery, background worker, or automatic retry supplied by this design.')
qa('Why so many forwarding classes?', 'Handlers and organisers establish workflow boundaries, but several currently just forward calls. I would retain a layer when it owns meaningful coordination or an assessment requirement, and simplify it when it adds no clear value.')
source('Application/Behaviors; Application/Vehicles/Commands; Application/Vehicles/Organisers; Application/Vehicles/Results/UpsertVehicleResult.cs; API/Controllers, all under backend/src/LevelMovement.*.')

page('Domain modelling and validation')
h('Entities and value objects')
p('Vehicle is an entity because it has a Guid identity that persists when its make, model, or rate changes. RegistrationNumber, VehicleYear, and DailyRate are sealed records with getter-only values. They are value objects: equality is based on their values rather than a separate identity.')
p('RegistrationNumber removes all whitespace, uppercases letters, rejects blank values, and limits the compact result to 20 characters. KW00KBGP and KW 00 KB GP are the same value. Punctuation is preserved. The frontend displays matching GP plates in pairs, while API values remain compact. VehicleYear accepts 1000 through 9999, not just realistic manufacturing years.')
p('DailyRate uses decimal and rounds to two decimal places with midpoint rounding to even. 10.125 becomes 10.12; 10.135 becomes 10.14. A value that rounds to zero or below is rejected. Both the value object and application validator enforce the maximum supported by decimal(18,2).')
qa('Why not use strings and decimals everywhere?', 'A named type communicates meaning and keeps normalisation and rules in one place. A RegistrationNumber cannot accidentally be passed where a VehicleYear is expected. The cost is extra mapping and a few more small types.')
h('Validation at three boundaries')
p('Frontend Zod validation gives fast feedback and reduces avoidable requests. Application FluentValidation protects the HTTP workflow and returns field-level errors. Domain constructors protect value rules when those objects are created outside the usual request path. The database adds durable constraints such as uniqueness, maximum column lengths, and numeric precision.')
p('These layers serve different purposes. Client validation cannot be trusted because callers can bypass the browser. An application-level “does it exist?” check cannot replace a unique index because two requests can pass the check at the same time.')
h('Encapsulation and entity behaviour')
p('Vehicle has private setters and a private constructor for EF. Register creates a valid, identified vehicle. UpdateDetails validates every new value before changing any field, and preserves identity, registration, and creation date. Delete marks an active vehicle as deleted. Deleted vehicles reject further edits and repeated domain deletion.')
p('EnsureEligibleForBooking belongs to Vehicle because it depends only on that vehicle’s state. Application BookingEligibilityService loads the entity and invokes the rule. No separate domain manager is needed for this single-entity rule. There is no date availability, reservation, payment, or booking record creation.')
source('backend/src/LevelMovement.Domain/ValueObjects/Vehicles; Domain/Vehicles/Vehicle.cs; Application/Services/BookingEligibilityService.cs; Application/Vehicles/Validators; frontend/src/features/vehicles/validation/vehicleSchema.ts.')

page('Persistence and concurrency')
h('Repository and unit of work')
p('VehicleRepository describes retrieval and changes to stored vehicles. UnitOfWork delegates SaveChangesAsync to the shared DbContext. Repository methods stage changes; the command behaviour commits them. EF Core already provides repository-like and unit-of-work-like facilities, so these wrappers add an application boundary rather than a new persistence engine.')
p('A single relational SaveChanges operation normally uses a transaction for its database changes. The wrapper does not automatically cover multiple saves, external HTTP calls, email, or a distributed workflow. Those need explicitly designed consistency boundaries.')
h('Tracking conversion and database constraints')
p('Tracked entities let EF detect property changes and compare original concurrency values. GetAllAsync uses AsNoTracking because its entities are only read. The details lookup currently uses a tracked query even though its response is read-only. VehicleConfig converts the value objects into string, integer, and decimal columns, keeping domain types without separate value-object tables.')
p('The unique registration index prevents duplicate registrations, including registrations on soft-deleted rows. CreatedDate is set with UTC time on creation. During updates, Vehicle.UpdateDetails changes editable fields and preserves identity and creation date. EF can materialize private setters and the interceptor can update the stamp through the change tracker.')
h('Optimistic concurrency in a two-user example')
for t in [
'Alice and Bob both load a vehicle with stamp A.',
'Alice submits stamp A. EF uses that supplied stamp as the original concurrency value. The interceptor assigns a new current stamp B, and the update succeeds.',
'Bob submits stamp A after Alice saves. The database row now has B, so the update using A matches no row. EF throws DbUpdateConcurrencyException, which becomes HTTP 409.',
'Bob must reload the latest details and decide how to reapply his edits. Automatically retrying with B would defeat the protection.'
]: doc.add_paragraph(t,'List Number')
p('The repository explicitly marks Make as modified, even when submitted values are unchanged. This forces an update and stamp check for a no-op edit too. Deletion supplies the original stamp and changes IsDeleted, so stale deletes receive the same protection.')
qa('Why map the response after saving?', 'The stamp changes in the save interceptor. UpsertVehicleResult retains the tracked entity, and ToResponse runs after SaveChanges, so the caller receives the new stamp rather than the old one.')
source('backend/src/LevelMovement.Infrastructure/Repositories/VehicleRepository .cs; Persistence/Configurations/VehicleConfig.cs; Persistence/Interceptors/VehicleConcurrencyInterceptor.cs; Application/Behaviors/UnitOfWorkBehavior.cs.')

page('Soft deletion APIs and errors')
h('Soft deletion and query filters')
p('The application DeleteAsync operation loads the vehicle and calls Vehicle.Delete, which sets IsDeleted to true instead of removing the row. VehicleConfig applies a global filter that hides deleted vehicles from normal queries. This preserves the record while keeping ordinary lists and details free of deleted items. It is not a full audit history: there is no complete record of who changed each field and when.')
p('GetForBookingAsync deliberately uses IgnoreQueryFilters so the service can distinguish a deleted vehicle from one that never existed. A deleted vehicle produces a business conflict; an unknown registration produces not found. Bypassing filters should be deliberate because it exposes records hidden from ordinary workflows.')
h('HTTP contracts to remember')
for q,a in [
('GET /api/vehicles','Returns 200 and active vehicle summaries. List items omit the concurrency stamp and internal entity ID.'),
('GET /api/vehicles/{registrationNumber}','Returns details and a stamp, or 404 when no visible vehicle exists.'),
('POST /api/vehicles','An omitted stamp means create; a non-empty stamp means update. Returns 201 or 200. An explicitly empty Guid is invalid. This is an application-specific upsert contract, not an automatic insert-if-update-missing operation.'),
('DELETE /api/vehicles/{registrationNumber}','Takes a JSON body with the stamp and returns 204 on success. The frontend correctly avoids parsing JSON from an empty 204 response.'),
('POST /api/bookings/validate','Returns an eligibility result without saving or creating a booking.')
]: qa(q,a)
h('Centralised exception handling')
p('ApiExceptionHandler maps FluentValidation failures to 400, missing entities to 404, and stale writes, unavailable vehicles, or SQL Server duplicate-key errors to 409. Unexpected failures become 500. ProblemDetails provides a consistent error shape; validation responses additionally contain field errors.')
p('Server logs include the exception, trace identifier, endpoint, request type, and failing function. Client responses use controlled messages and remove traceId. RequestLoggingBehavior supplies request context; it does not itself measure duration or log every successful request.')
qa('What is the tradeoff in hiding trace IDs?', 'The response stays minimal, but users have no correlation ID to quote to support. An opaque correlation ID can be useful without exposing stack traces or secrets, if that fits the product requirements.')
source('backend/src/LevelMovement.Api/Controllers; Middleware/ApiExceptionHandler.cs; Behaviors/RequestLoggingBehavior.cs; Infrastructure/Repositories/VehicleRepository .cs; frontend/src/shared/api/client.ts.')

page('React TypeScript and testing')
h('Views and controller hooks')
p('VehicleListPage and VehicleEditPage render the UI. Hooks such as useVehicleEditController coordinate forms, fetching, saving, deleting, and navigation. vehiclesApi defines endpoint calls, and the shared request function handles HTTP and error parsing. This is a separation similar to Model View Controller, adapted to React; the hooks are not ASP.NET controllers.')
p('React state holds values that affect rendering, such as a success message or delete confirmation. Refs retain mutable values without causing a render. formStamp is a ref tied to the form snapshot, so a background query-cache refresh cannot silently replace the concurrency stamp for the user’s older edits.')
h('Server state and form state')
p('TanStack Query manages fetched data, loading, errors, and cached results. React Hook Form manages the editable input snapshot. Saving updates the cached detail and invalidates the list; editing also refetches canonical details. Deleting removes the detail cache and invalidates the list. Explicit reload resets the form and its stamp after a conflict.')
p('This distinction matters: the latest cache data and a user’s unsaved form are not always the same version. Automatically resetting the form on every fetch could destroy edits. The hook initializes the form once and uses an explicit reload path for conflicts.')
qa('Does TypeScript validate JSON?', 'No. TypeScript types are compile-time checks. The shared request function parses JSON; vehiclesApi now validates vehicle and booking responses with Zod schemas. Form input has its own schema. Runtime validation protects the UI from malformed response shapes.')
h('What the tests demonstrate')
p('VehicleDomainTests checks valid creation, unchanged state after rejected edits, immutable identity, and deleted-vehicle rules. VehicleValueObjectTests checks rounding, equality, registration normalisation, year boundaries, and EF mapping metadata. VehicleApiTests and VehicleCrudTests exercise HTTP behaviour through WebApplicationFactory with SQLite, including stale updates and deletes, soft deletion, field errors, and response stamps. These are integration-style tests even though they live under a Unit folder.')
p('Frontend tests use Testing Library and mocked API interactions to exercise creation, duplicate registrations, conflicts, stamp preservation, deletion, and field focus. A test with mocked networking proves UI behaviour against the mock; it does not prove the deployed frontend and SQL Server work together.')
qa('What would you test next?', 'SQL Server duplicate-key translation and migration deployment against the actual provider, plus a browser workflow against the real API. Tests should target observable behaviour and risks, not merely repeat implementation details.')
source('frontend/src/features/vehicles/controllers; frontend/src/shared/api/client.ts; frontend/src/features/vehicles/vehicles.test.tsx; backend/tests/Unit/LevelMovement.Tests.')

page('Tradeoffs and improvements to discuss')
h('Be accurate about what is implemented')
p('Say “layered architecture with a Vehicle aggregate that protects its invariants, value objects, repositories, and application services.” DDD is a modelling approach rather than a compliance badge. Do not claim microservices, event sourcing, durable messaging, or a complete booking system. Say “optimistic concurrency” rather than “database locking.” Say “soft deletion” rather than “complete audit logging.”')
h('Reduce layers where they add no behaviour')
p('Several handlers and organisers remain forwarding wrappers. Workflow coordination and response mapping now live in Application; Infrastructure focuses on persistence. UpsertVehicleOrganiser adapts the compatibility endpoint into explicit operations. Further simplification should remove only layers that add no useful boundary.')
h('Make rules consistent')
p('The browser requires a rate of at least 0.01, while the backend first rounds and can accept a smaller positive input that rounds to 0.01. The browser trims make and model, while backend validation and persistence do not apply the same trimming. Shared business examples and boundary tests would make these differences explicit.')
p('JavaScript number uses binary floating point and cannot preserve every large amount accepted by a backend decimal(18,2). A smaller business maximum, decimal strings, or a carefully specified minor-unit contract could improve precision. Choose a representation based on the actual currency and magnitude requirements.')
h('Evolve the external contract')
p('The entity now owns its transitions and DailyRate enforces its upper bound. The external contract still has room to improve: stable machine-readable error codes would be more robust than matching an English duplicate-registration title in the frontend.')
p('The current delete endpoint carries a JSON body. That works with the project client, but a version precondition such as If-Match may be easier for some HTTP tooling. Separate create and update endpoints could also make intent clearer than a stamp-driven upsert.')
h('Plan for growth and deployment')
p('The list currently loads all visible vehicles without pagination. Server-side filtering, bounded pages, DTO projection, and query-plan inspection would be appropriate as the dataset grows. Add authentication and operation-level authorisation before exposing protected data; no such configuration appears in the current Program.cs.')
p('Docker Compose starts SQL Server, waits for health, runs a migration job, and then starts the API. A named volume preserves data across container replacement. This provides a reproducible local environment, but production rollout still needs deliberate secrets, migration, backup, monitoring, and recovery plans.')
source('backend/compose.yaml; backend/src/LevelMovement.Api/Program.cs; Application/Services; Domain/Vehicles/Vehicle.cs; frontend/src/features/vehicles/validation/vehicleSchema.ts and controllers/useVehicleEditController.ts.')

page('DDD changes and how to explain them')
h('The Vehicle aggregate protects its rules')
p('An aggregate is a consistency boundary, not a required base class. Here Vehicle and its value objects form a small aggregate. Its public business methods are the entry points for changes. Application code cannot assign identity, registration, creation date, deletion state, or editable fields directly. EF is a trusted persistence mechanism and can reconstruct saved state through private members.')
qa('What changed?', 'Previously, Vehicle was a mutable data container and the repository validated and changed its fields. Now Vehicle.Register, UpdateDetails, Delete, and EnsureEligibleForBooking enforce the rules. The repository adds or stages the resulting entity and supplies the original concurrency value.')
h('Creation and editing are distinct operations')
p('RegisterAsync creates a fresh entity with a new ID and UTC creation date. UpdateAsync retrieves the existing entity, then invokes UpdateDetails. It never constructs a temporary Vehicle with an empty ID. A missing update target still returns 404; it does not become an insert. The existing HTTP upsert shape is preserved, with the stamp interpreted by the organiser at the application boundary.')
h('Domain rules and application coordination')
p('The application service chooses the sequence: load, call the domain method, stage persistence. The entity decides which state changes are valid. BookingEligibilityService performs a read and calls the entity rule. The old VehicleBookingManager was removed because this rule fits on one entity. A domain service can be introduced later if a business rule genuinely spans concepts and does not fit an entity or value object.')
h('Database consistency remains essential')
p('Private setters do not prevent two separate requests from loading the same version. The repository still supplies the expected stamp, EF checks it, and the save interceptor generates the next stamp. The database unique index still arbitrates competing registrations. Entity invariants, concurrency checks, and uniqueness constraints protect different boundaries.')
h('What this refactor does not claim')
p('This is a focused DDD refactor for vehicle management. It does not implement a reservation aggregate, availability by date, domain events, or separate bounded contexts. Those require real business requirements. There is no need to add them just to collect pattern names. No schema change is required for private setters and service relocation; the earlier registration-normalisation migration remains a separate data migration.')
qa('An answer to practise', '“I moved business invariants into Vehicle and exposed named operations instead of public setters. Application services coordinate the use cases and Infrastructure handles persistence. I preserved the API contract and tested identity preservation, rejected edits, soft deletion, and stale writes.”')
source('backend/src/LevelMovement.Domain/Vehicles/Vehicle.cs; backend/src/LevelMovement.Application/Services; backend/src/LevelMovement.Infrastructure/Repositories/VehicleRepository .cs; backend/tests/Unit/LevelMovement.Tests/VehicleDomainTests.cs.')

page('Intermediate concepts you should know')
qa('Async and await', 'Async I/O allows a request thread to avoid blocking while a database operation waits. It does not necessarily create a new thread or make the database query faster. Avoid .Result and .Wait() in request flows. The project passes cancellation tokens through the backend; cancellation is cooperative, not a guarantee of rollback after a commit.')
qa('Generics and constraints', 'A generic type works with different types while retaining compile-time checks. UnitOfWorkBehavior<TRequest,TResponse> constrains TRequest to ICommand<TResponse>, which expresses that saving belongs to commands. Generic constraints communicate capabilities required by the implementation.')
qa('Records classes and nullability', 'Records provide value-oriented equality; ordinary classes normally use reference equality unless overridden. A record is not automatically deeply immutable. Nullable reference annotations and required properties improve compile-time checking, but do not replace runtime business validation.')
qa('LINQ and query execution', 'IQueryable describes a provider-translated query; operators before ToListAsync can become SQL. IEnumerable operates over an enumerable sequence, often in memory. The list repository filters and orders before materialisation, but the service maps entities to DTOs after loading them. Projection before materialisation could reduce transferred columns.')
qa('Transactions and race conditions', 'A transaction groups related database changes so they succeed or fail together. Isolation affects what concurrent transactions observe. Optimistic concurrency detects an outdated version; a unique constraint rejects duplicates. Neither replaces the other, and a prior read is not protection against a later race.')
qa('HTTP safety and idempotency', 'Safe operations do not request a state change. Idempotent operations have the same intended final effect when repeated. A repeated delete can return a different status while leaving the same deleted state. The current create POST is not a general idempotency mechanism: retrying a successful create can produce a duplicate conflict.')
qa('Authentication authorisation and CORS', 'Authentication establishes identity; authorisation decides permitted actions. CORS controls browser access across origins and is not an authorisation mechanism. A caller outside a browser can still invoke an unprotected API.')
qa('Testing boundaries', 'Unit tests isolate a small rule; integration tests check components together; end-to-end tests exercise a complete user workflow. Mocking an interface is useful, but too many mocks can hide contract errors. Use the real database provider for provider-specific behaviour.')
qa('Performance and observability', 'Measure before optimising. Understand indexes, query plans, pagination, N+1 queries, and memory allocations. Logs explain events, metrics reveal trends, and traces connect operations. Logging exceptions is helpful, but it is not a complete monitoring system.')
qa('Git and delivery', 'Be able to explain a branch, commit, pull request, conflict resolution, and rollback. A useful CI pipeline restores dependencies, builds, and runs relevant tests. Database migrations are versioned schema changes and should be reviewed alongside application changes.')

page('Practice questions and revision')
qa('Why use a repository when EF already exists?', 'It gives the application a persistence contract and a place for vehicle-specific queries. I recognise the extra indirection, and would avoid a generic wrapper that only duplicates DbSet without adding value.')
qa('Why validate twice?', 'Client validation improves usability; backend validation protects the system from every caller. Domain types protect their own invariants. The rules should agree where their responsibilities overlap.')
qa('Can two users overwrite each other?', 'They can attempt to, but the original concurrency stamp is checked during the write. The second stale update fails with 409 and must reload instead of silently winning.')
qa('Why is registration uniqueness a database rule?', 'Only a database constraint arbitrates concurrent inserts reliably. A pre-check improves messages but cannot remove the race between checking and inserting.')
qa('Does a successful eligibility check reserve the vehicle?', 'No. It only confirms the current deletion rule. State can change after the check. Real booking would need an atomic reservation operation and rules for overlapping dates.')
qa('Is everything here SOLID?', 'There are strong examples of single responsibility and dependency inversion. SOLID is a set of design principles, not a certification. I would explain the protected entity operations and discuss remaining forwarding layers and contract boundaries as tradeoffs.')
qa('What happens if saving fails?', 'The command does not return a successful HTTP result because the save behaviour throws before the controller maps the response. The exception handler then translates the failure. Changes to other systems would require separate coordination.')
h('Exercises you should be able to do')
for t in [
'Draw the save flow from the React form to SQL Server. Point to where validation, stamping, saving, and DTO mapping happen.',
'Simulate two clients with one original stamp. Predict the responses for the first save, second save, and stale delete.',
'Explain why a deleted registration remains unavailable for reuse under the current unique index.',
'Write a test case for a rate that rounds to zero and a rate exactly halfway between two cents.',
'Describe how you would add pagination without loading the full table into memory.',
'Explain one design you would simplify and one rule you would strengthen, including the cost of the change.'
]: doc.add_paragraph(t,'List Bullet')
h('A strong answer structure')
p('“The problem was ____. In this project I used ____ in ____. It works by ____. That gives us ____. The tradeoff is ____, and I would change it if ____.” Use a concrete example from the code instead of listing pattern names. If a feature is planned rather than implemented, say so.')

out=Path('C:/Development/LevelMovement/docs/LevelMovement_Principles_Study_Guide.docx')
doc.save(out)
markdown = []
for paragraph in doc.paragraphs:
    if not paragraph.text.strip():
        continue
    style = paragraph.style.name
    prefix = '# ' if style == 'Title' else '## ' if style == 'Heading 1' else '### ' if style == 'Heading 2' else '- ' if style == 'List Bullet' else '1. ' if style == 'List Number' else ''
    markdown.append(prefix + paragraph.text)
out.with_suffix('.md').write_text('\n\n'.join(markdown) + '\n', encoding='utf-8')
print(out)
